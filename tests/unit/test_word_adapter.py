"""Unit tests for Microsoft Word native conversion adapter and detection.

Tests:
- WordInfo defaults and candidate detection.
- App Paths registry detection mocking.
- WordAdapter conversion capability and plan generation.
- DOCX -> PDF lossless conversion via Word COM automation.
- Validation checks via validate_pdf.
"""

from __future__ import annotations

import sys
from pathlib import Path
from unittest.mock import patch

import pytest

from cuttle.adapters.word import WordAdapter
from cuttle.engine.errors import CuttleError
from cuttle.engine.types import CancelToken, Quality
from cuttle.platform.windows.word_detect import (
    WordInfo,
    detect_word,
)


def test_word_info_defaults() -> None:
    info = WordInfo(installed=False)
    assert not info.installed
    assert info.path is None
    assert info.version is None
    assert info.name == "Microsoft Word"


def test_detect_word_fake_directory(tmp_path: Path) -> None:
    fake_office_dir = tmp_path / "Microsoft Office" / "root" / "Office16"
    fake_office_dir.mkdir(parents=True)
    fake_winword = fake_office_dir / "WINWORD.EXE"
    fake_winword.write_bytes(b"MZfakebinary")

    info = detect_word(
        search_paths=[fake_winword],
        check_registry=False,
    )
    assert info.installed is True
    assert info.path == fake_winword


def test_detect_word_registry_mock(tmp_path: Path) -> None:
    fake_exe = tmp_path / "WINWORD.EXE"
    fake_exe.write_bytes(b"fake")

    with patch(
        "cuttle.platform.windows.word_detect._query_registry_app_path",
        return_value=fake_exe,
    ):
        info = detect_word(
            search_paths=[],
            check_registry=True,
        )

    if sys.platform == "win32":
        assert info.installed is True
        assert info.path == fake_exe


def test_word_adapter_can_convert() -> None:
    adapter = WordAdapter()
    assert adapter.can_convert("docx", "pdf") == Quality.LOSSLESS
    assert adapter.can_convert("docx", "png") is None
    assert adapter.can_convert("pdf", "docx") is None
    assert adapter.can_convert("xlsx", "pdf") is None


def test_word_adapter_plan(tmp_path: Path) -> None:
    adapter = WordAdapter()
    src = tmp_path / "doc.docx"
    tmp_out = tmp_path / "doc.pdf.cuttle-part"

    plan = adapter.plan(src, "pdf", tmp_out, {})
    assert plan.args == ["word", str(src), "pdf", str(tmp_out)]

    with pytest.raises(CuttleError) as exc_info:
        adapter.plan(src, "docx", tmp_out, {})
    assert exc_info.value.code == "E_UNSUPPORTED"


def test_word_adapter_run_mocked(tmp_path: Path) -> None:
    adapter = WordAdapter()
    src = tmp_path / "sample.docx"
    src.write_bytes(b"PK0304mock")
    tmp_out = tmp_path / "out.pdf"

    plan = adapter.plan(src, "pdf", tmp_out, {})
    cancel = CancelToken()

    with patch("cuttle.platform.windows.word_com.convert_docx_to_pdf_word") as mock_com:
        mock_com.return_value = None
        adapter.run(plan, lambda p, msg: None, cancel)
        mock_com.assert_called_once_with(src, tmp_out, timeout_s=60.0)


def test_word_adapter_run_cancelled(tmp_path: Path) -> None:
    adapter = WordAdapter()
    src = tmp_path / "sample.docx"
    tmp_out = tmp_path / "out.pdf"
    plan = adapter.plan(src, "pdf", tmp_out, {})

    cancel = CancelToken()
    cancel.cancel()

    with pytest.raises(CuttleError) as exc_info:
        adapter.run(plan, lambda p, msg: None, cancel)
    assert exc_info.value.code == "E_CANCELLED"


@pytest.mark.skipif(sys.platform != "win32", reason="Windows COM only")
def test_word_adapter_real_if_installed(tmp_path: Path) -> None:
    """Test actual Word conversion if Word is installed on the host machine."""
    info = detect_word()
    if not info.installed:
        pytest.skip("Microsoft Word not installed on host machine.")

    import docx

    src_docx = tmp_path / "hello_word.docx"
    doc = docx.Document()
    doc.add_heading("Hello World from Cuttle", level=0)
    p = doc.add_paragraph("High-fidelity native layout preservation test.")
    p.add_run(" Bold text.").bold = True
    doc.save(str(src_docx))

    tmp_out = tmp_path / "hello_word.pdf"
    adapter = WordAdapter()
    plan = adapter.plan(src_docx, "pdf", tmp_out, {})
    cancel = CancelToken()

    adapter.run(plan, lambda p, msg: None, cancel)

    assert tmp_out.is_file()
    assert tmp_out.stat().st_size > 0

    val = adapter.validate(src_docx, tmp_out)
    assert val.ok
